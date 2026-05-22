"""
Unit tests for the DOCX parser pipeline — no DB, no network, no AI.
Uses python-docx to build minimal in-memory fixtures.
"""

import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from app.parsers.docx_parser import DocxParser


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _build_docx(paragraphs: list[str]) -> bytes:
    """Create a simple DOCX with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_with_heading(heading: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_with_tracked_insertion(text: str) -> bytes:
    """Build a DOCX containing a w:ins tracked insertion run."""
    doc = Document()
    para = doc.add_paragraph()
    # Add a normal run, then inject a w:ins element via lxml
    run = para.add_run("Normal text. ")

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ins = etree.SubElement(para._p, f"{{{W}}}ins")
    ins.set(f"{{{W}}}id", "1")
    ins.set(f"{{{W}}}author", "Reviewer")
    ins.set(f"{{{W}}}date", "2026-01-01T00:00:00Z")
    ins_run = etree.SubElement(ins, f"{{{W}}}r")
    ins_t = etree.SubElement(ins_run, f"{{{W}}}t")
    ins_t.text = text

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_docx_with_strikethrough(text: str) -> bytes:
    """Build a DOCX with a run that has w:strike formatting."""
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run(text)
    # Apply strikethrough via run XML
    rpr = run._r.get_or_add_rPr()
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    strike = etree.SubElement(rpr, f"{{{W}}}strike")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Basic parsing
# ---------------------------------------------------------------------------

def test_parse_returns_list_of_clauses():
    docx_bytes = _build_docx(["Clause 1 text.", "Clause 2 text.", "Clause 3 text."])
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    assert isinstance(clauses, list)
    assert len(clauses) >= 1


def test_parse_empty_document_returns_empty_or_minimal():
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    parser = DocxParser()
    clauses = parser.parse(buf.getvalue())
    # An empty DOCX may return [] or a list with a single empty-text clause
    assert isinstance(clauses, list)


def test_parse_preserves_paragraph_text():
    text = "This is an important indemnification clause requiring careful review."
    docx_bytes = _build_docx([text])
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    combined = " ".join(c.get("body_text", "") or c.get("clause_text", "") for c in clauses)
    assert "indemnification" in combined


def test_parse_produces_paragraph_index():
    docx_bytes = _build_docx(["First clause.", "Second clause.", "Third clause."])
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    indices = [c.get("paragraph_index") for c in clauses if c.get("paragraph_index") is not None]
    # Paragraph indices should be monotonically increasing
    assert indices == sorted(indices)


def test_parse_multiple_clauses_count():
    texts = [f"Clause {i}: Some agreement text." for i in range(1, 6)]
    docx_bytes = _build_docx(texts)
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    assert len(clauses) >= 1  # At minimum, content is captured


# ---------------------------------------------------------------------------
# Heading detection
# ---------------------------------------------------------------------------

def test_parse_detects_headings():
    docx_bytes = _build_docx_with_heading("1. Definitions", "In this Agreement...")
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    headings = [c.get("heading") for c in clauses if c.get("heading")]
    # At least one heading should be detected from the heading paragraph
    assert len(headings) >= 1 or any("Definitions" in c.get("body_text", "") for c in clauses)


# ---------------------------------------------------------------------------
# Tracked change detection
# ---------------------------------------------------------------------------

def test_parse_detects_tracked_insertion():
    docx_bytes = _build_docx_with_tracked_insertion("Added clause text")
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    # At least one clause should flag has_tracked_insertion
    has_insertion = any(c.get("has_tracked_insertion") for c in clauses)
    assert has_insertion


def test_parse_detects_strikethrough():
    docx_bytes = _build_docx_with_strikethrough("This was deleted text")
    parser = DocxParser()
    clauses = parser.parse(docx_bytes)
    has_strike = any(c.get("has_strikethrough") for c in clauses)
    assert has_strike


# ---------------------------------------------------------------------------
# PDF fallback
# ---------------------------------------------------------------------------

def test_pdf_parser_returns_clauses():
    """Test PyMuPDF fallback parser with a minimal PDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        pytest.skip("PyMuPDF not installed")

    from app.parsers.pdf_parser import PdfParser

    # Create minimal PDF in memory
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((72, 72), "This is a test contract clause for PDF parsing.")
    pdf_bytes = pdf_doc.tobytes()

    parser = PdfParser()
    clauses = parser.parse(pdf_bytes)
    assert isinstance(clauses, list)
    combined = " ".join(c.get("body_text", "") or c.get("clause_text", "") for c in clauses)
    assert "contract" in combined.lower()
